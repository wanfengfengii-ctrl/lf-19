import os
import io
from datetime import datetime
from typing import List, Dict, Optional

from database import DatabaseManager, Building
from services.deviation_calculator import DeviationCalculator, ComponentDeviation


class ReportGenerator:
    def __init__(self, db_manager: DatabaseManager):
        self.db = db_manager
        self.calculator = DeviationCalculator(db_manager)

    def generate_text_report(self, building_id: int) -> str:
        building = self.db.get_building(building_id)
        if not building:
            return "建筑不存在"

        deviations = self.calculator.calculate_building_deviations(building_id)
        stats = self.calculator.get_deviation_statistics(building_id)
        abnormal_list = self.calculator.get_abnormal_components(building_id)
        recheck_stats = self.db.get_recheck_statistics(building_id)

        lines = []
        lines.append("=" * 60)
        lines.append("古建筑木构件尺寸偏差检测报告")
        lines.append("=" * 60)
        lines.append(f"建筑名称：{building.name}")
        lines.append(f"建筑编号：{building.code}")
        lines.append(f"建筑位置：{building.location}")
        lines.append(f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("")

        lines.append("-" * 60)
        lines.append("一、总体统计")
        lines.append("-" * 60)
        lines.append(f"构件总数：{stats['total_components']} 件")
        lines.append(f"已检测构件：{stats['has_data_count']} 件")
        lines.append(f"未检测构件：{stats['no_data_count']} 件")
        lines.append(f"合格构件：{stats['normal_count']} 件")
        lines.append(f"异常构件：{stats['abnormal_count']} 件")
        if stats['total_components'] > 0:
            lines.append(f"合格率：{stats['normal_count'] / stats['total_components'] * 100:.1f}%")
        lines.append(f"待复测任务：{recheck_stats['pending']} 件")
        lines.append(f"复测中：{recheck_stats['in_progress']} 件")
        lines.append(f"已完成复测：{recheck_stats['completed']} 件")
        lines.append("")

        lines.append("-" * 60)
        lines.append("二、偏差统计")
        lines.append("-" * 60)
        len_stat = stats['length_stats']
        lines.append(f"长度偏差 - 最小:{len_stat['min']:.2f}  最大:{len_stat['max']:.2f}  平均:{len_stat['avg']:.2f}")
        wid_stat = stats['width_stats']
        lines.append(f"宽度偏差 - 最小:{wid_stat['min']:.2f}  最大:{wid_stat['max']:.2f}  平均:{wid_stat['avg']:.2f}")
        ang_stat = stats['angle_stats']
        lines.append(f"角度偏差 - 最小:{ang_stat['min']:.2f}  最大:{ang_stat['max']:.2f}  平均:{ang_stat['avg']:.2f}")
        lines.append("")

        lines.append("-" * 60)
        lines.append("三、异常构件明细")
        lines.append("-" * 60)
        if abnormal_list:
            for i, dev in enumerate(abnormal_list, 1):
                lines.append(f"{i}. 构件编号：{dev.component.code}")
                lines.append(f"   构件名称：{dev.component.name or '-'}")
                lines.append(f"   构件类型：{dev.component.component_type or '-'}")
                if dev.length_deviation and dev.length_deviation.is_abnormal:
                    lines.append(f"   长度偏差：设计值 {dev.length_deviation.design_value:.2f}，"
                                 f"实测值 {dev.length_deviation.measured_value:.2f}，"
                                 f"偏差 {dev.length_deviation.deviation:+.2f} "
                                 f"({dev.length_deviation.deviation_percent:.1f}%) [异常]")
                if dev.width_deviation and dev.width_deviation.is_abnormal:
                    lines.append(f"   宽度偏差：设计值 {dev.width_deviation.design_value:.2f}，"
                                 f"实测值 {dev.width_deviation.measured_value:.2f}，"
                                 f"偏差 {dev.width_deviation.deviation:+.2f} "
                                 f"({dev.width_deviation.deviation_percent:.1f}%) [异常]")
                if dev.angle_deviation and dev.angle_deviation.is_abnormal:
                    lines.append(f"   角度偏差：设计值 {dev.angle_deviation.design_value:.2f}°，"
                                 f"实测值 {dev.angle_deviation.measured_value:.2f}°，"
                                 f"偏差 {dev.angle_deviation.deviation:+.2f}° "
                                 f"({dev.angle_deviation.deviation_percent:.1f}%) [异常]")
                if dev.recheck_needed:
                    lines.append(f"   状态：需复测")
                lines.append("")
        else:
            lines.append("无异常构件")
            lines.append("")

        lines.append("-" * 60)
        lines.append("四、全构件检测明细")
        lines.append("-" * 60)
        lines.append(f"{'序号':<4} {'构件编号':<12} {'构件名称':<12} {'长度偏差':<12} {'宽度偏差':<12} {'角度偏差':<12} {'状态':<8}")
        lines.append("-" * 76)
        for i, dev in enumerate(deviations, 1):
            ld = f"{dev.length_deviation.deviation:+.2f}" if dev.length_deviation else "-"
            wd = f"{dev.width_deviation.deviation:+.2f}" if dev.width_deviation else "-"
            ad = f"{dev.angle_deviation.deviation:+.2f}°" if dev.angle_deviation else "-"
            status = "异常" if dev.is_abnormal else "正常"
            if not dev.latest_record:
                status = "未测"
            elif dev.recheck_needed:
                status = "需复测"
            lines.append(f"{i:<4} {dev.component.code:<12} {dev.component.name or '-':<12} "
                         f"{ld:<12} {wd:<12} {ad:<12} {status:<8}")

        lines.append("")
        lines.append("=" * 60)
        lines.append("报告结束")
        lines.append("=" * 60)

        return "\n".join(lines)

    def save_report_to_file(self, building_id: int, file_path: str) -> bool:
        try:
            report_text = self.generate_text_report(building_id)
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(report_text)
            return True
        except Exception as e:
            print(f"保存报告失败：{e}")
            return False

    def generate_excel_report(self, building_id: int, file_path: str) -> bool:
        try:
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                from openpyxl.utils import get_column_letter
            except ImportError:
                return False

            building = self.db.get_building(building_id)
            if not building:
                return False

            deviations = self.calculator.calculate_building_deviations(building_id)
            stats = self.calculator.get_deviation_statistics(building_id)
            abnormal_list = self.calculator.get_abnormal_components(building_id)
            recheck_stats = self.db.get_recheck_statistics(building_id)

            wb = Workbook()

            header_font = Font(bold=True, size=12)
            title_font = Font(bold=True, size=14)
            normal_font = Font(size=11)
            red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            center_align = Alignment(horizontal='center', vertical='center')

            ws1 = wb.active
            ws1.title = "总体概览"

            ws1.merge_cells('A1:F1')
            cell = ws1['A1']
            cell.value = "古建筑木构件尺寸偏差检测报告"
            cell.font = title_font
            cell.alignment = center_align

            ws1['A3'] = "建筑名称"
            ws1['B3'] = building.name
            ws1['A4'] = "建筑编号"
            ws1['B4'] = building.code
            ws1['A5'] = "建筑位置"
            ws1['B5'] = building.location
            ws1['A6'] = "报告生成时间"
            ws1['B6'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

            for row in ws1.iter_rows(min_row=3, max_row=6, min_col=1, max_col=1):
                for cell in row:
                    cell.font = header_font

            ws1['A8'] = "统计项目"
            ws1['B8'] = "数量"
            ws1['A8'].font = header_font
            ws1['B8'].font = header_font
            ws1['A8'].fill = gray_fill
            ws1['B8'].fill = gray_fill

            overview_data = [
                ("构件总数", stats['total_components']),
                ("已检测构件", stats['has_data_count']),
                ("未检测构件", stats['no_data_count']),
                ("合格构件", stats['normal_count']),
                ("异常构件", stats['abnormal_count']),
                ("待复测任务", recheck_stats['pending']),
                ("复测中", recheck_stats['in_progress']),
                ("已完成复测", recheck_stats['completed']),
            ]
            for i, (name, value) in enumerate(overview_data, start=9):
                ws1[f'A{i}'] = name
                ws1[f'B{i}'] = value

            ws1.column_dimensions['A'].width = 20
            ws1.column_dimensions['B'].width = 20

            ws2 = wb.create_sheet("偏差统计")
            ws2['A1'] = "维度"
            ws2['B1'] = "最小偏差"
            ws2['C1'] = "最大偏差"
            ws2['D1'] = "平均偏差"
            for col in ['A', 'B', 'C', 'D']:
                ws2[f'{col}1'].font = header_font
                ws2[f'{col}1'].fill = gray_fill
                ws2[f'{col}1'].alignment = center_align

            len_stat = stats['length_stats']
            wid_stat = stats['width_stats']
            ang_stat = stats['angle_stats']

            ws2['A2'] = "长度"
            ws2['B2'] = round(len_stat['min'], 2)
            ws2['C2'] = round(len_stat['max'], 2)
            ws2['D2'] = round(len_stat['avg'], 2)

            ws2['A3'] = "宽度"
            ws2['B3'] = round(wid_stat['min'], 2)
            ws2['C3'] = round(wid_stat['max'], 2)
            ws2['D3'] = round(wid_stat['avg'], 2)

            ws2['A4'] = "角度"
            ws2['B4'] = round(ang_stat['min'], 2)
            ws2['C4'] = round(ang_stat['max'], 2)
            ws2['D4'] = round(ang_stat['avg'], 2)

            for col in ['A', 'B', 'C', 'D']:
                ws2.column_dimensions[col].width = 15

            ws3 = wb.create_sheet("全构件明细")
            headers = ["序号", "构件编号", "构件名称", "构件类型",
                       "设计长度", "实测长度", "长度偏差",
                       "设计宽度", "实测宽度", "宽度偏差",
                       "设计角度", "实测角度", "角度偏差",
                       "状态"]
            for i, h in enumerate(headers, start=1):
                cell = ws3.cell(row=1, column=i, value=h)
                cell.font = header_font
                cell.fill = gray_fill
                cell.alignment = center_align
                cell.border = thin_border

            for i, dev in enumerate(deviations, start=2):
                row_data = [
                    i - 1,
                    dev.component.code,
                    dev.component.name or "-",
                    dev.component.component_type or "-",
                ]

                if dev.length_deviation:
                    row_data.extend([
                        dev.length_deviation.design_value,
                        dev.length_deviation.measured_value,
                        round(dev.length_deviation.deviation, 2),
                    ])
                else:
                    row_data.extend(["-", "-", "-"])

                if dev.width_deviation:
                    row_data.extend([
                        dev.width_deviation.design_value,
                        dev.width_deviation.measured_value,
                        round(dev.width_deviation.deviation, 2),
                    ])
                else:
                    row_data.extend(["-", "-", "-"])

                if dev.angle_deviation:
                    row_data.extend([
                        dev.angle_deviation.design_value,
                        dev.angle_deviation.measured_value,
                        round(dev.angle_deviation.deviation, 2),
                    ])
                else:
                    row_data.extend(["-", "-", "-"])

                if not dev.latest_record:
                    status = "未测"
                elif dev.is_abnormal:
                    status = "异常"
                elif dev.recheck_needed:
                    status = "需复测"
                else:
                    status = "正常"
                row_data.append(status)

                for j, val in enumerate(row_data, start=1):
                    cell = ws3.cell(row=i, column=j, value=val)
                    cell.border = thin_border
                    cell.alignment = center_align

                    if status == "异常":
                        cell.fill = red_fill
                    elif status == "正常":
                        cell.fill = green_fill
                    elif status == "需复测":
                        cell.fill = yellow_fill

            for col_idx in range(1, len(headers) + 1):
                ws3.column_dimensions[get_column_letter(col_idx)].width = 14

            ws4 = wb.create_sheet("异常构件明细")
            abnormal_headers = ["序号", "构件编号", "构件名称", "异常维度",
                                "设计值", "实测值", "偏差值", "偏差百分比", "阈值"]
            for i, h in enumerate(abnormal_headers, start=1):
                cell = ws4.cell(row=1, column=i, value=h)
                cell.font = header_font
                cell.fill = gray_fill
                cell.alignment = center_align
                cell.border = thin_border

            row_idx = 2
            for i, dev in enumerate(abnormal_list, 1):
                if dev.length_deviation and dev.length_deviation.is_abnormal:
                    self._add_abnormal_row(ws4, row_idx, i, dev, "长度",
                                           dev.length_deviation)
                    row_idx += 1

                if dev.width_deviation and dev.width_deviation.is_abnormal:
                    self._add_abnormal_row(ws4, row_idx, i, dev, "宽度",
                                           dev.width_deviation)
                    row_idx += 1

                if dev.angle_deviation and dev.angle_deviation.is_abnormal:
                    self._add_abnormal_row(ws4, row_idx, i, dev, "角度",
                                           dev.angle_deviation)
                    row_idx += 1

            for col_idx in range(1, len(abnormal_headers) + 1):
                ws4.column_dimensions[get_column_letter(col_idx)].width = 14

            wb.save(file_path)
            return True

        except Exception as e:
            print(f"生成Excel报告失败：{e}")
            return False

    def _add_abnormal_row(self, ws, row_idx, seq, dev, dim_name, dev_item):
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center')

        data = [
            seq,
            dev.component.code,
            dev.component.name or "-",
            dim_name,
            dev_item.design_value,
            dev_item.measured_value,
            round(dev_item.deviation, 2),
            f"{dev_item.deviation_percent:.2f}%",
            dev_item.threshold,
        ]
        for j, val in enumerate(data, start=1):
            cell = ws.cell(row=row_idx, column=j, value=val)
            cell.fill = red_fill
            cell.border = thin_border
            cell.alignment = center_align

    def generate_pdf_report(self, building_id: int, file_path: str) -> bool:
        try:
            try:
                from reportlab.lib.pagesizes import A4
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.lib.units import mm
                from reportlab.lib import colors
                from reportlab.platypus import (
                    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
                )
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont
            except ImportError:
                return False

            font_registered = False
            font_paths = [
                "/System/Library/Fonts/PingFang.ttc",
                "/System/Library/Fonts/STHeiti Light.ttc",
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            ]
            for fp in font_paths:
                if os.path.exists(fp):
                    try:
                        pdfmetrics.registerFont(TTFont('CustomFont', fp))
                        font_registered = True
                        break
                    except Exception:
                        continue

            font_name = 'CustomFont' if font_registered else 'Helvetica'

            building = self.db.get_building(building_id)
            if not building:
                return False

            deviations = self.calculator.calculate_building_deviations(building_id)
            stats = self.calculator.get_deviation_statistics(building_id)
            abnormal_list = self.calculator.get_abnormal_components(building_id)
            recheck_stats = self.db.get_recheck_statistics(building_id)

            doc = SimpleDocTemplate(file_path, pagesize=A4,
                                    leftMargin=20*mm, rightMargin=20*mm,
                                    topMargin=20*mm, bottomMargin=20*mm)

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle', parent=styles['Title'],
                fontName=font_name, fontSize=18, leading=22,
                spaceAfter=10*mm, textColor=colors.black
            )
            heading_style = ParagraphStyle(
                'CustomHeading', parent=styles['Heading2'],
                fontName=font_name, fontSize=14, leading=18,
                spaceBefore=6*mm, spaceAfter=3*mm
            )
            normal_style = ParagraphStyle(
                'CustomNormal', parent=styles['Normal'],
                fontName=font_name, fontSize=10, leading=14
            )
            small_style = ParagraphStyle(
                'CustomSmall', parent=styles['Normal'],
                fontName=font_name, fontSize=9, leading=12
            )

            story = []

            story.append(Paragraph("古建筑木构件尺寸偏差检测报告", title_style))
            story.append(Paragraph(f"建筑名称：{building.name}", normal_style))
            story.append(Paragraph(f"建筑编号：{building.code}", normal_style))
            story.append(Paragraph(f"建筑位置：{building.location}", normal_style))
            story.append(Paragraph(f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", normal_style))
            story.append(Spacer(1, 8*mm))

            story.append(Paragraph("一、总体统计", heading_style))
            overview_data = [
                ["统计项目", "数量"],
                ["构件总数", str(stats['total_components'])],
                ["已检测构件", str(stats['has_data_count'])],
                ["未检测构件", str(stats['no_data_count'])],
                ["合格构件", str(stats['normal_count'])],
                ["异常构件", str(stats['abnormal_count'])],
                ["待复测任务", str(recheck_stats['pending'])],
                ["复测中", str(recheck_stats['in_progress'])],
                ["已完成复测", str(recheck_stats['completed'])],
            ]
            overview_table = Table(overview_data, colWidths=[60*mm, 40*mm])
            overview_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), font_name),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(overview_table)

            story.append(Paragraph("二、偏差统计", heading_style))
            len_stat = stats['length_stats']
            wid_stat = stats['width_stats']
            ang_stat = stats['angle_stats']

            dev_stats_data = [
                ["维度", "最小偏差", "最大偏差", "平均偏差"],
                ["长度", f"{len_stat['min']:.2f}", f"{len_stat['max']:.2f}", f"{len_stat['avg']:.2f}"],
                ["宽度", f"{wid_stat['min']:.2f}", f"{wid_stat['max']:.2f}", f"{wid_stat['avg']:.2f}"],
                ["角度", f"{ang_stat['min']:.2f}", f"{ang_stat['max']:.2f}", f"{ang_stat['avg']:.2f}"],
            ]
            dev_table = Table(dev_stats_data, colWidths=[30*mm, 30*mm, 30*mm, 30*mm])
            dev_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), font_name),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(dev_table)

            story.append(PageBreak())

            story.append(Paragraph("三、异常构件明细", heading_style))
            if abnormal_list:
                for i, dev in enumerate(abnormal_list, 1):
                    story.append(Paragraph(
                        f"{i}. 构件编号：{dev.component.code}（{dev.component.name or '未命名'}）",
                        normal_style
                    ))
                    if dev.length_deviation and dev.length_deviation.is_abnormal:
                        story.append(Paragraph(
                            f"&nbsp;&nbsp;&nbsp;&nbsp;长度偏差：设计 {dev.length_deviation.design_value:.2f}，"
                            f"实测 {dev.length_deviation.measured_value:.2f}，"
                            f"偏差 {dev.length_deviation.deviation:+.2f} "
                            f"({dev.length_deviation.deviation_percent:.1f}%) 【异常】",
                            small_style
                        ))
                    if dev.width_deviation and dev.width_deviation.is_abnormal:
                        story.append(Paragraph(
                            f"&nbsp;&nbsp;&nbsp;&nbsp;宽度偏差：设计 {dev.width_deviation.design_value:.2f}，"
                            f"实测 {dev.width_deviation.measured_value:.2f}，"
                            f"偏差 {dev.width_deviation.deviation:+.2f} "
                            f"({dev.width_deviation.deviation_percent:.1f}%) 【异常】",
                            small_style
                        ))
                    if dev.angle_deviation and dev.angle_deviation.is_abnormal:
                        story.append(Paragraph(
                            f"&nbsp;&nbsp;&nbsp;&nbsp;角度偏差：设计 {dev.angle_deviation.design_value:.2f}°，"
                            f"实测 {dev.angle_deviation.measured_value:.2f}°，"
                            f"偏差 {dev.angle_deviation.deviation:+.2f}° "
                            f"({dev.angle_deviation.deviation_percent:.1f}%) 【异常】",
                            small_style
                        ))
                    if dev.recheck_needed:
                        story.append(Paragraph("&nbsp;&nbsp;&nbsp;&nbsp;状态：需复测", small_style))
                    story.append(Spacer(1, 3*mm))
            else:
                story.append(Paragraph("无异常构件", normal_style))

            story.append(PageBreak())

            story.append(Paragraph("四、全构件检测明细", heading_style))
            detail_data = [["序号", "构件编号", "构件名称", "长度偏差", "宽度偏差", "角度偏差", "状态"]]
            for i, dev in enumerate(deviations, 1):
                ld = f"{dev.length_deviation.deviation:+.2f}" if dev.length_deviation else "-"
                wd = f"{dev.width_deviation.deviation:+.2f}" if dev.width_deviation else "-"
                ad = f"{dev.angle_deviation.deviation:+.2f}°" if dev.angle_deviation else "-"
                if not dev.latest_record:
                    status = "未测"
                elif dev.is_abnormal:
                    status = "异常"
                elif dev.recheck_needed:
                    status = "需复测"
                else:
                    status = "正常"
                detail_data.append([
                    str(i), dev.component.code, dev.component.name or "-",
                    ld, wd, ad, status
                ])

            detail_table = Table(detail_data, colWidths=[10*mm, 20*mm, 20*mm, 18*mm, 18*mm, 18*mm, 16*mm])
            table_style = [
                ('FONTNAME', (0, 0), (-1, -1), font_name),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]

            for i in range(1, len(detail_data)):
                status = detail_data[i][-1]
                if status == "异常":
                    table_style.append(('BACKGROUND', (0, i), (-1, i), colors.Color(1, 0.8, 0.8)))
                elif status == "正常":
                    table_style.append(('BACKGROUND', (0, i), (-1, i), colors.Color(0.8, 1, 0.8)))
                elif status == "需复测":
                    table_style.append(('BACKGROUND', (0, i), (-1, i), colors.Color(1, 1, 0.8)))

            detail_table.setStyle(TableStyle(table_style))
            story.append(detail_table)

            doc.build(story)
            return True

        except Exception as e:
            print(f"生成PDF报告失败：{e}")
            import traceback
            traceback.print_exc()
            return False
